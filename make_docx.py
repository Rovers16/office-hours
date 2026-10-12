#!/usr/bin/env python3
"""Turn the CSV written by the C program into a Word document.

All scheduling decisions are made by the C program; this script only lays the
results out in the style of the sample tables (a CLASS SCHEDULE table and an
OFFICE HOURS list for each professor).

Usage:
    python3 make_docx.py output/office_hours.csv output/office_hours.docx
"""

import csv
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
NOTE = "NOTE:  Instructors are not to be interrupted during class times."


def read_report(path):
    """Group the CSV rows by professor, keeping the order they appear in."""
    profs = []
    by_name = {}
    with open(path, newline="", encoding="utf-8") as f:
        for row in csv.reader(f):
            if not row:
                continue
            kind, name = row[0], row[1]
            if kind == "PROF":
                prof = {"name": name, "department": row[2], "office": row[3],
                        "email": row[4], "source": row[5], "classes": [], "hours": []}
                profs.append(prof)
                by_name[name] = prof
            elif kind == "CLASS":
                by_name[name]["classes"].append(row[2:5])
            elif kind == "OFFICE_HOUR":
                by_name[name]["hours"].append(row[2:5])
            else:
                raise ValueError(f"{path}: unknown record type {kind!r}")
    return profs


def add_professor(doc, prof):
    doc.add_heading(prof["name"], level=1)
    details = [prof["department"], f"Office: {prof['office']}" if prof["office"] else "",
               prof["email"]]
    doc.add_paragraph("   |   ".join(d for d in details if d))

    doc.add_heading("CLASS SCHEDULE", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Course", "Times", "Location"]):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for course, times, room in prof["classes"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = course, times, room

    source = "provided by the lecturer" if prof["source"] == "provided" else "generated"
    doc.add_heading(f"OFFICE HOURS ({source})", level=2)
    for day in DAYS:
        blocks = [(t.replace(" - ", " – "), c) for d, t, c in prof["hours"] if d == day]
        if not blocks and day in ("Saturday", "Sunday"):
            continue
        text = "\n\t& ".join(f"{t} ({c})" if c else t for t, c in blocks)
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(24)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.tab_stops.add_tab_stop(Inches(1.4))
        para.add_run(f"{day}:").bold = True
        para.add_run(f"\t{text}")

    doc.add_paragraph().add_run(NOTE).italic = True


def main(argv):
    if len(argv) != 3:
        print(__doc__.strip().splitlines()[-1].strip(), file=sys.stderr)
        return 2
    csv_path, docx_path = Path(argv[1]), Path(argv[2])

    doc = Document()
    doc.add_heading("Teaching and Office-Hours Schedule", level=0)
    profs = read_report(csv_path)
    for i, prof in enumerate(profs):
        if i:
            doc.add_page_break()
        add_professor(doc, prof)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"wrote {docx_path} ({len(profs)} professors)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
